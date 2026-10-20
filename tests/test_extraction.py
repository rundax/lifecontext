import dataclasses
import struct
import unittest
from hashlib import sha256
from io import BytesIO
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from lifecontext.extraction import (
    ExtractionLimits,
    Locator,
    TextSegment,
    _ocr_image,
    _Rejected,
    extract_document,
)


def office(entries=None, compression=0):
    buffer = BytesIO()
    with ZipFile(buffer, "w", compression=compression) as archive:
        for name, content in (
            entries
            or {"[Content_Types].xml": b"<Types/>", "word/document.xml": b"<document/>"}
        ).items():
            archive.writestr(name, content)
    return buffer.getvalue()


class ExtractionTests(unittest.TestCase):
    def extract(self, data=b"hello\nworld", filename="note.txt", **kwargs):
        return extract_document(data, filename, source_id="document-1", **kwargs)

    def test_text_provenance_and_ordered_locators(self):
        raw = b"\xef\xbb\xbf# Heading\r\n\nUnicode: \xc3\xa9"
        artifact = self.extract(raw, "note.md", source_version=2, extraction_revision=3)
        self.assertEqual(artifact.status, "completed")
        self.assertEqual(artifact.original_checksum, sha256(raw).hexdigest())
        self.assertEqual(
            (artifact.source_version, artifact.extraction_revision), (2, 3)
        )
        self.assertEqual(artifact.text, "# Heading\r\n\nUnicode: é")
        self.assertEqual([s.locator.start for s in artifact.segments], [1, 2, 3])
        self.assertEqual(artifact.segments[-1].end, len(artifact.text))
        self.assertEqual(
            artifact.artifact_checksum,
            self.extract(
                raw, "note.md", source_version=2, extraction_revision=3
            ).artifact_checksum,
        )
        self.assertNotEqual(
            artifact.artifact_checksum,
            dataclasses.replace(artifact, extraction_revision=4).artifact_checksum,
        )

    def test_text_rejections(self):
        for data, reason in [
            (b"\xff", "invalid_utf8"),
            (b" \n", "no_extractable_text"),
            (b"a\x00b", "text_signature_mismatch"),
            (b"%PDF-1.7", "text_signature_mismatch"),
        ]:
            with self.subTest(reason=reason):
                result = self.extract(data)
                self.assertEqual(result.reason, reason)
                self.assertFalse(result.segments)

    def test_filename_checks(self):
        for filename in [
            "../a.txt",
            "a\\b.txt",
            "a:stream.txt",
            "a\x00.txt",
            "",
            "a.txt.",
            "a" * 256,
        ]:
            with self.subTest(filename=filename):
                self.assertEqual(
                    self.extract(filename=filename).reason, "unsafe_filename"
                )

    def test_independent_text_limits(self):
        for limits, reason in [
            (ExtractionLimits(max_original_bytes=2), "original_byte_limit"),
            (ExtractionLimits(max_text_bytes=2), "text_byte_limit"),
            (ExtractionLimits(max_text_characters=2), "text_character_limit"),
        ]:
            self.assertEqual(self.extract(limits=limits).reason, reason)
        self.assertEqual(
            self.extract(
                "é".encode(), limits=ExtractionLimits(max_text_characters=1)
            ).status,
            "completed",
        )

    def test_limits_validation_and_identity(self):
        for field in dataclasses.fields(ExtractionLimits):
            for bad in [0, -1, True, float("inf"), float("nan")]:
                with (
                    self.subTest(field=field.name, value=bad),
                    self.assertRaises(ValueError),
                ):
                    ExtractionLimits(**{field.name: bad})
        self.assertNotEqual(
            ExtractionLimits().identity, ExtractionLimits(max_pixels=1).identity
        )

    def test_binary_backends_are_blocked_not_empty_success(self):
        for filename, data in [
            ("a.pdf", b"%PDF-1.7\n"),
            ("a.png", b"\x89PNG\r\n\x1a\n"),
            ("a.jpg", b"\xff\xd8\xff"),
            ("a.tiff", b"II*\x00"),
            ("a.docx", office()),
        ]:
            with self.subTest(filename=filename):
                result = self.extract(data, filename)
                self.assertEqual(result.status, "failed")
                self.assertFalse(result.segments)
        self.assertEqual(self.extract(b"wrong", "a.pdf").reason, "signature_mismatch")
        self.assertEqual(self.extract(filename="a.doc").status, "unsupported")
        self.assertEqual(
            self.extract(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1", "a.docx").status, "locked"
        )

    def test_office_format_contracts(self):
        for ext, part in [
            ("xlsx", "xl/workbook.xml"),
            ("pptx", "ppt/presentation.xml"),
        ]:
            self.assertEqual(
                self.extract(
                    office({"[Content_Types].xml": b"<Types/>", part: b"<root/>"}),
                    "a." + ext,
                ).status,
                "failed",
            )
        self.assertEqual(
            self.extract(office(), "a.xlsx").reason, "ooxml_signature_mismatch"
        )

    def test_zip_limits(self):
        raw = office()
        for limits in [
            ExtractionLimits(max_zip_entries=1),
            ExtractionLimits(max_zip_entry_bytes=1),
            ExtractionLimits(max_zip_expanded_bytes=1),
        ]:
            self.assertEqual(
                self.extract(raw, "a.docx", limits=limits).status, "limit_exceeded"
            )
        bomb = office(
            {"word/document.xml": b"a" * 10000, "[Content_Types].xml": b"<Types/>"},
            ZIP_DEFLATED,
        )
        self.assertEqual(
            self.extract(
                bomb, "a.docx", limits=ExtractionLimits(max_zip_ratio=2)
            ).reason,
            "zip_expansion_limit",
        )

    def test_zip_paths_macros_malformed_and_crc(self):
        for path in ["../evil", "/evil", "a\\evil", "a//evil", "a/./evil", "a:evil"]:
            self.assertEqual(
                self.extract(office({path: b"x"}), "a.docx").reason, "unsafe_zip_path"
            )
        raw = office(
            {
                "[Content_Types].xml": b"<Types/>",
                "word/document.xml": b"<document/>",
                "word/vbaProject.bin": b"never executed",
            }
        )
        self.assertEqual(self.extract(raw, "a.docx").status, "unsupported")
        self.assertEqual(
            self.extract(b"PK\x03\x04broken", "a.docx").reason, "malformed_archive"
        )
        raw = bytearray(office())
        name_length, extra_length = struct.unpack_from("<HH", raw, 26)
        raw[30 + name_length + extra_length] ^= 1
        self.assertEqual(self.extract(bytes(raw), "a.docx").reason, "malformed_archive")

    def test_encrypted_zip(self):
        raw = bytearray(office())
        struct.pack_into("<H", raw, 6, 1)
        central = raw.index(b"PK\x01\x02")
        struct.pack_into("<H", raw, central + 8, 1)
        self.assertEqual(self.extract(bytes(raw), "a.docx").status, "locked")

    def test_cooperative_timeout(self):
        with patch("lifecontext.extraction.time.monotonic", side_effect=[0, ExtractionLimits().max_seconds + 1]):
            self.assertEqual(self.extract().reason, "time_limit")
        with patch("lifecontext.extraction.time.monotonic", side_effect=[0, ExtractionLimits().max_seconds + 1]):
            self.assertEqual(self.extract(office(), "a.docx").reason, "time_limit")

    def test_contract_validation(self):
        for args in [
            ("lines", 0, 1),
            ("page", 2, 1),
            ("unknown", 1, 1),
            ("sheet_cell", 1, 1),
        ]:
            with self.assertRaises(ValueError):
                Locator(*args)
        with self.assertRaises(ValueError):
            TextSegment("abc", 0, 2, Locator("lines", 1, 1))
        artifact = self.extract()
        for changes in [
            {"segments": ()},
            {"source_version": True},
            {"original_checksum": "bad"},
            {"segments": (artifact.segments[1],)},
            {"status": "blocked"},
            {"segments": list(artifact.segments)},
        ]:
            with self.subTest(changes=changes), self.assertRaises(ValueError):
                dataclasses.replace(artifact, **changes)


class BinaryExtractionTests(unittest.TestCase):
    extract = ExtractionTests.extract

    def test_docx_order_tables_and_nested_cells(self):
        from docx import Document

        document = Document()
        document.add_paragraph("Before")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Cell one"
        table.cell(0, 1).text = "Cell two"
        table.cell(0, 1).add_table(rows=1, cols=1).cell(0, 0).text = "Nested"
        document.add_paragraph("After")
        buffer = BytesIO()
        document.save(buffer)
        artifact = self.extract(buffer.getvalue(), "a.docx")
        self.assertEqual(artifact.status, "completed", artifact.reason)
        self.assertEqual(artifact.text, "Before\nCell one\nCell two\nNested\nAfter\n")
        self.assertEqual(artifact.segments[1].locator.kind, "table_cell")
        self.assertIn("table:", artifact.segments[3].locator.container)
        self.assertEqual(
            self.extract(
                buffer.getvalue(), "a.docx", limits=ExtractionLimits(max_cells=1)
            ).reason,
            "cell_limit",
        )

    def test_xlsx_formula_cache_sheet_and_cell_limits(self):
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active.title = "Budget"
        workbook.active.append(["Salary", 123, "=B1*2"])
        workbook.create_sheet("Other").append(["End"])
        buffer = BytesIO()
        workbook.save(buffer)
        artifact = self.extract(buffer.getvalue(), "a.xlsx")
        self.assertEqual(artifact.status, "completed", artifact.reason)
        self.assertEqual(artifact.text, "Salary\n123\nEnd\n")
        self.assertEqual(artifact.segments[1].locator.container, "Budget!B1")
        self.assertIn("formula_cached", artifact.warnings[0])
        self.assertEqual(
            self.extract(
                buffer.getvalue(), "a.xlsx", limits=ExtractionLimits(max_sheets=1)
            ).reason,
            "sheet_limit",
        )
        self.assertEqual(
            self.extract(
                buffer.getvalue(), "a.xlsx", limits=ExtractionLimits(max_cells=1)
            ).reason,
            "cell_limit",
        )

    def test_pptx_group_table_order_and_slide_limit(self):
        from pptx import Presentation
        from pptx.util import Inches

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        group = slide.shapes.add_group_shape()
        group.shapes.add_textbox(0, 0, Inches(1), Inches(1)).text = "Grouped"
        table = slide.shapes.add_table(1, 1, 0, 0, Inches(1), Inches(1)).table
        table.cell(0, 0).text = "Table"
        presentation.slides.add_slide(presentation.slide_layouts[6])
        buffer = BytesIO()
        presentation.save(buffer)
        artifact = self.extract(buffer.getvalue(), "a.pptx")
        self.assertEqual(artifact.status, "completed", artifact.reason)
        self.assertEqual(artifact.text, "Grouped\nTable\n")
        self.assertEqual(artifact.segments[0].locator.kind, "slide")
        self.assertEqual(artifact.segments[1].locator.kind, "table_cell")
        self.assertEqual(
            self.extract(
                buffer.getvalue(), "a.pptx", limits=ExtractionLimits(max_slides=1)
            ).reason,
            "slide_limit",
        )

    def pdf(self, text=True, encrypted=False, pages=1, drawing=False):
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

        writer = PdfWriter()
        for _ in range(pages):
            page = writer.add_blank_page(width=100, height=100)
            if drawing:
                # Visible ink without a text layer: a scan-like page that needs OCR.
                stream = DecodedStreamObject()
                stream.set_data(b"0 g 10 10 80 80 re f")
                page[NameObject("/Contents")] = writer._add_object(stream)
            elif text:
                font = DictionaryObject(
                    {
                        NameObject("/Type"): NameObject("/Font"),
                        NameObject("/Subtype"): NameObject("/Type1"),
                        NameObject("/BaseFont"): NameObject("/Helvetica"),
                    }
                )
                page[NameObject("/Resources")] = DictionaryObject(
                    {
                        NameObject("/Font"): DictionaryObject(
                            {NameObject("/F1"): writer._add_object(font)}
                        )
                    }
                )
                stream = DecodedStreamObject()
                stream.set_data(b"BT /F1 12 Tf 10 50 Td (Hello PDF) Tj ET")
                page[NameObject("/Contents")] = writer._add_object(stream)
        if encrypted:
            writer.encrypt("secret")
        buffer = BytesIO()
        writer.write(buffer)
        return buffer.getvalue()

    def test_book_length_pdf_preserves_all_page_locators(self):
        artifact = self.extract(self.pdf(pages=5000), "book.pdf")
        self.assertEqual(artifact.status, "completed", artifact.reason)
        self.assertEqual(len(artifact.segments), 5000)
        self.assertEqual([s.locator.start for s in artifact.segments], list(range(1, 5001)))
        self.assertEqual(artifact.text.count("Hello PDF"), 5000)

    def test_pdf_above_default_page_limit_is_rejected(self):
        artifact = self.extract(self.pdf(pages=5001), "oversized.pdf")
        self.assertEqual(artifact.status, "limit_exceeded")
        self.assertEqual(artifact.reason, "page_limit")

    def test_text_pdf_encryption_page_limits_and_scanned_block(self):
        artifact = self.extract(self.pdf(pages=2), "a.pdf")
        self.assertEqual(artifact.status, "completed", artifact.reason)
        self.assertEqual(artifact.text, "Hello PDF\nHello PDF\n")
        self.assertEqual([s.locator.start for s in artifact.segments], [1, 2])
        self.assertEqual(
            self.extract(self.pdf(encrypted=True), "a.pdf").status, "locked"
        )
        self.assertEqual(
            self.extract(
                self.pdf(pages=2), "a.pdf", limits=ExtractionLimits(max_pages=1)
            ).reason,
            "page_limit",
        )
        # Blank detection runs before OCR, so a blank page never needs Tesseract.
        self.assertEqual(
            self.extract(self.pdf(text=False), "a.pdf").reason,
            "no_extractable_text",
        )
        scanned = self.extract(self.pdf(drawing=True), "a.pdf")
        self.assertEqual(
            (scanned.status, scanned.reason), ("blocked", "ocr_backend_not_configured")
        )

    def test_images_pixel_frame_limits_and_missing_ocr(self):
        from PIL import Image, ImageDraw

        image = Image.new("RGB", (10, 10), "white")
        inked = Image.new("RGB", (10, 10), "white")
        ImageDraw.Draw(inked).rectangle((2, 2, 7, 7), fill="black")
        for format, name in [("PNG", "a.png"), ("JPEG", "a.jpg"), ("TIFF", "a.tiff")]:
            buffer = BytesIO()
            image.save(buffer, format=format)
            self.assertEqual(
                self.extract(buffer.getvalue(), name).reason,
                "no_extractable_text",
            )
            drawn = BytesIO()
            inked.save(drawn, format=format)
            self.assertEqual(
                self.extract(drawn.getvalue(), name).reason,
                "ocr_backend_not_configured",
            )
            self.assertEqual(
                self.extract(
                    buffer.getvalue(), name, limits=ExtractionLimits(max_pixels=10)
                ).reason,
                "pixel_limit",
            )
        buffer = BytesIO()
        image.save(buffer, format="TIFF", save_all=True, append_images=[image.copy()])
        self.assertEqual(
            self.extract(
                buffer.getvalue(), "a.tiff", limits=ExtractionLimits(max_frames=1)
            ).reason,
            "frame_limit",
        )

    def test_ocr_fixed_argv_output_cleanup_and_timeout(self):
        import subprocess
        import tempfile
        from pathlib import Path

        from PIL import Image

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            image = Image.new("RGB", (10, 10))

            def success(argv, **kwargs):
                self.assertEqual(argv[0], "/usr/bin/tesseract")
                self.assertNotIn("shell", kwargs)
                (root / "ocr.txt").write_text("Recognized")
                return subprocess.CompletedProcess(argv, 0)

            with patch("lifecontext.extraction._ocr_languages", return_value="eng"), patch("lifecontext.extraction.subprocess.run", side_effect=success):
                self.assertEqual(
                    _ocr_image(image, root, ExtractionLimits(), "/usr/bin/tesseract"),
                    "Recognized",
                )
            self.assertFalse((root / "ocr.png").exists())
            with patch(
                "lifecontext.extraction.subprocess.run",
                side_effect=subprocess.TimeoutExpired("ocr", 1),
            ):
                with self.assertRaises(_Rejected) as caught:
                    _ocr_image(image, root, ExtractionLimits(), "/usr/bin/tesseract")
                self.assertEqual(caught.exception.reason, "ocr_time_limit")

    def test_png_growth_respects_remaining_aggregate_budget(self):
        import tempfile
        from pathlib import Path

        from PIL import Image

        image = Image.new("RGB", (32, 32), "white")
        encoded = BytesIO()
        image.save(encoded, format="PNG")
        png_size = len(encoded.getvalue())
        # PNG fits each individual file cap but not the aggregate remaining space.
        limits = ExtractionLimits(
            max_text_bytes=128,
            max_output_bytes=1024,
            max_temp_bytes=1000 + 128 + 1024 + png_size - 1,
        )
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "original").write_bytes(b"x" * 800)
            (root / "request.json").write_bytes(b"x" * 200)
            with patch("lifecontext.extraction.subprocess.run") as fake_ocr:
                with self.assertRaises(_Rejected) as caught:
                    _ocr_image(image, root, limits, "/mock/tesseract")
                self.assertEqual(caught.exception.reason, "temp_byte_limit")
                fake_ocr.assert_not_called()
            self.assertFalse((root / "ocr.png").exists())
            self.assertLessEqual(
                sum(p.stat().st_size for p in root.iterdir()), limits.max_temp_bytes
            )

    def test_budget_writer_checks_seeks_and_ocr_child_cap(self):
        import resource
        import subprocess
        import tempfile
        from pathlib import Path

        from PIL import Image

        from lifecontext.extraction import _BudgetWriter

        writer = _BudgetWriter(BytesIO(), 4)
        writer.write(b"1234")
        writer.seek(0)
        writer.write(b"a")
        with self.assertRaises(_Rejected):
            writer.seek(5)
        writer.seek(4)
        with self.assertRaises(_Rejected):
            writer.write(b"a")
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)

            def fake_ocr(argv, **kwargs):
                with patch("resource.setrlimit") as setter:
                    kwargs["preexec_fn"]()
                    self.assertEqual(setter.call_args.args[0], resource.RLIMIT_FSIZE)
                    self.assertLessEqual(setter.call_args.args[1][0], 128)
                (root / "ocr.txt").write_text("OCR")
                return subprocess.CompletedProcess(argv, 0)

            with patch("lifecontext.extraction._ocr_languages", return_value="eng"), patch("lifecontext.extraction.subprocess.run", side_effect=fake_ocr):
                result = _ocr_image(
                    Image.new("RGB", (10, 10)),
                    root,
                    ExtractionLimits(max_text_bytes=128),
                    "/mock/tesseract",
                )
                self.assertEqual(result, "OCR")

    def test_request_bytes_counted_before_worker_launch(self):
        with patch("lifecontext.extraction.subprocess.Popen") as launch:
            artifact = self.extract(
                self.pdf(),
                "a.pdf",
                limits=ExtractionLimits(
                    max_output_bytes=128, max_temp_bytes=len(self.pdf()) + 128
                ),
            )
            self.assertEqual(artifact.reason, "temp_byte_limit")
            launch.assert_not_called()

    def test_worker_timeout_and_output_limits(self):
        self.assertEqual(
            self.extract(
                self.pdf(), "a.pdf", limits=ExtractionLimits(max_seconds=0.001)
            ).status,
            "limit_exceeded",
        )
        self.assertEqual(
            self.extract(
                self.pdf(), "a.pdf", limits=ExtractionLimits(max_text_characters=2)
            ).reason,
            "text_limit",
        )
        self.assertEqual(
            self.extract(
                self.pdf(), "a.pdf", limits=ExtractionLimits(max_temp_bytes=1)
            ).reason,
            "temp_byte_limit",
        )
        self.assertEqual(
            self.extract(self.pdf(), "a.pdf", tesseract_path="relative").reason,
            "ocr_executable_unavailable",
        )

    def test_directory_cap_precedes_zipfile_and_zlib_errors(self):
        raw = office()
        with patch("lifecontext.extraction.ZipFile") as constructor:
            self.assertEqual(
                self.extract(
                    raw, "a.docx", limits=ExtractionLimits(max_zip_entries=1)
                ).reason,
                "zip_entry_limit",
            )
            constructor.assert_not_called()
        raw = bytearray(office(compression=ZIP_DEFLATED))
        name_length, extra_length = struct.unpack_from("<HH", raw, 26)
        raw[30 + name_length + extra_length] = 255
        self.assertEqual(self.extract(bytes(raw), "a.docx").reason, "malformed_archive")

    def test_xml_entities_rejected(self):
        payload = b'<!DOCTYPE root [<!ENTITY x "evil">]><root>&x;</root>'
        self.assertEqual(
            self.extract(
                office(
                    {"[Content_Types].xml": b"<Types/>", "word/document.xml": payload}
                ),
                "a.docx",
            ).status,
            "failed",
        )

    def test_mixed_pdf_keeps_text_and_marks_unresolved_ocr_page(self):
        import tempfile
        from pathlib import Path

        from pypdf import PdfReader, PdfWriter

        from lifecontext.extraction import _binary_extract
        writer = PdfWriter()
        writer.add_page(PdfReader(BytesIO(self.pdf())).pages[0])
        writer.add_page(PdfReader(BytesIO(self.pdf(text=False))).pages[0])
        writer.add_page(PdfReader(BytesIO(self.pdf(drawing=True))).pages[0])
        stream = BytesIO()
        writer.write(stream)
        with tempfile.TemporaryDirectory() as directory, patch("lifecontext.extraction._ocr_image", return_value=""):
            result = _binary_extract(stream.getvalue(), "pdf", ExtractionLimits(), Path(directory), "/mock/tesseract")
            self.assertEqual(result["status"], "completed")
            self.assertEqual([s["locator"]["start"] for s in result["segments"]], [1])
            self.assertEqual(
                result["warnings"],
                [
                    "blank_page:2; no visible content",
                    "ocr_no_text_page:3; page content unresolved",
                ],
            )
            # Nothing read: the failure keeps its ledger instead of an empty one.
            failed = _binary_extract(self.pdf(drawing=True), "pdf", ExtractionLimits(), Path(directory), "/mock/tesseract")
            self.assertEqual((failed["status"], failed["reason"]), ("failed", "no_extractable_text"))
            self.assertEqual(failed["warnings"], ["ocr_no_text_page:1; page content unresolved"])

    def mixed_pdf(self):
        from pypdf import PdfReader, PdfWriter

        writer = PdfWriter()
        writer.add_page(PdfReader(BytesIO(self.pdf())).pages[0])
        writer.add_page(PdfReader(BytesIO(self.pdf(drawing=True))).pages[0])
        stream = BytesIO()
        writer.write(stream)
        return stream.getvalue()

    def test_single_page_ocr_failures_never_fail_the_document(self):
        import tempfile
        from pathlib import Path

        from lifecontext.extraction import _binary_extract

        for reason in ("ocr_time_limit", "ocr_failed", "ocr_output_limit"):
            with self.subTest(reason=reason), tempfile.TemporaryDirectory() as directory:
                with patch(
                    "lifecontext.extraction._ocr_image",
                    side_effect=_Rejected("limit_exceeded", reason),
                ):
                    result = _binary_extract(
                        self.mixed_pdf(), "pdf", ExtractionLimits(), Path(directory), "/mock/tesseract"
                    )
                self.assertEqual(result["status"], "completed")
                self.assertEqual(
                    result["warnings"],
                    [f"ocr_failed_page:2; {reason}; page content unresolved"],
                )
        # A document-wide budget still stops the whole extraction.
        with tempfile.TemporaryDirectory() as directory, patch(
            "lifecontext.extraction._ocr_image",
            side_effect=_Rejected("limit_exceeded", "temp_byte_limit"),
        ):
            with self.assertRaises(_Rejected) as caught:
                _binary_extract(
                    self.mixed_pdf(), "pdf", ExtractionLimits(), Path(directory), "/mock/tesseract"
                )
            self.assertEqual(caught.exception.reason, "temp_byte_limit")

    def test_mixed_pdf_without_tesseract_keeps_text_and_marks_page(self):
        artifact = self.extract(self.mixed_pdf(), "a.pdf")
        self.assertEqual(artifact.status, "completed", artifact.reason)
        self.assertEqual(artifact.text, "Hello PDF\n")
        self.assertEqual(
            artifact.warnings, ("ocr_unavailable_page:2; page content unresolved",)
        )

    def test_figures_on_text_pages_are_flagged(self):
        from PIL import Image
        from pypdf import PdfReader, PdfWriter

        picture = BytesIO()
        Image.new("RGB", (100, 100), "gray").save(picture, format="PDF")
        writer = PdfWriter(clone_from=BytesIO(picture.getvalue()))
        writer.pages[0].merge_page(PdfReader(BytesIO(self.pdf())).pages[0])
        writer.add_page(PdfReader(BytesIO(self.pdf())).pages[0])
        stream = BytesIO()
        writer.write(stream)
        artifact = self.extract(stream.getvalue(), "figure.pdf")
        self.assertEqual(artifact.status, "completed", artifact.reason)
        self.assertEqual(artifact.text, "Hello PDF\nHello PDF\n")
        self.assertEqual(
            artifact.warnings,
            ("figure_not_read:1; images or drawings on this page were not read",),
        )
        quiet = self.extract(
            stream.getvalue(), "figure.pdf", limits=ExtractionLimits(figure_min_area_bp=10_001)
        )
        self.assertEqual(quiet.warnings, ())

    def test_failed_documents_keep_their_ledger_and_are_never_complete(self):
        from lifecontext.extraction import coverage_summary

        artifact = self.extract(self.pdf(drawing=True, pages=2), "scan.pdf")
        self.assertEqual((artifact.status, artifact.reason), ("blocked", "ocr_backend_not_configured"))
        self.assertEqual(
            artifact.warnings,
            (
                "ocr_unavailable_page:1; page content unresolved",
                "ocr_unavailable_page:2; page content unresolved",
            ),
        )
        coverage = coverage_summary(artifact.to_dict())
        self.assertEqual((coverage["unresolved"], coverage["complete"]), (["page:1", "page:2"], False))
        self.assertFalse(coverage_summary({"format": "pdf", "segments": [], "warnings": []})["complete"])

    def test_ocr_stops_before_the_worker_deadline_and_keeps_the_rest(self):
        import tempfile
        import time
        from pathlib import Path

        from lifecontext.extraction import _binary_extract

        with tempfile.TemporaryDirectory() as directory, patch(
            "lifecontext.extraction._ocr_image", return_value="never"
        ) as ocr:
            result = _binary_extract(
                self.mixed_pdf(), "pdf", ExtractionLimits(), Path(directory), "/mock/tesseract",
                deadline=time.monotonic(),
            )
        ocr.assert_not_called()
        self.assertEqual(result["status"], "completed")
        self.assertEqual(result["warnings"], ["ocr_failed_page:2; time_limit; page content unresolved"])

    def overlay_pdf(self, content, *, picture_scale=1.0, pages=1):
        """Pages with a gray picture (full page, or scaled into a corner) under a text stream."""
        from PIL import Image
        from pypdf import PdfReader, PdfWriter, Transformation
        from pypdf.generic import DecodedStreamObject, NameObject

        picture = BytesIO()
        Image.new("RGB", (100, 100), "gray").save(picture, format="PDF")
        writer = PdfWriter()
        for _ in range(pages):
            text_page = PdfReader(BytesIO(self.pdf())).pages[0]
            stream = DecodedStreamObject()
            stream.set_data(content)
            text_page[NameObject("/Contents")] = stream
            page = writer.add_page(text_page)
            image_page = PdfReader(BytesIO(picture.getvalue())).pages[0]
            page.merge_transformed_page(
                image_page, Transformation().scale(picture_scale).translate(2, 2), over=False
            )
        buffer = BytesIO()
        writer.write(buffer)
        return buffer.getvalue()

    def test_searchable_scans_are_not_flagged_but_captioned_photos_are(self):
        words = b"Scanned words with an OCR layer " * 8
        scan = self.extract(self.overlay_pdf(b"BT /F1 2 Tf 3 Tr 1 50 Td (" + words + b") Tj ET"), "scan.pdf")
        self.assertEqual(scan.status, "completed", scan.reason)
        self.assertGreater(len(scan.text), 200)
        self.assertEqual(scan.warnings, ())
        photo = self.extract(self.overlay_pdf(b"BT /F1 2 Tf 1 50 Td (" + words + b") Tj ET"), "photo.pdf")
        self.assertEqual(photo.warnings, ("figure_not_read:1; images or drawings on this page were not read",))

    def test_repeated_small_logos_are_page_furniture(self):
        content = b"BT /F1 12 Tf 10 50 Td (Hello PDF) Tj ET"
        logos = self.extract(self.overlay_pdf(content, picture_scale=0.2, pages=3), "logos.pdf")
        self.assertEqual(logos.status, "completed", logos.reason)
        self.assertEqual(logos.warnings, ())
        twice = self.extract(self.overlay_pdf(content, picture_scale=0.2, pages=2), "twice.pdf")
        self.assertEqual(len(twice.warnings), 2)

    def test_wide_and_float_images_flatten_by_their_real_range(self):
        from PIL import Image

        from lifecontext.extraction import _flatten, _is_blank

        for mode, value in (("I", 65535), ("F", 0.25), ("F", 5000.0)):
            with self.subTest(mode=mode, value=value):
                uniform = Image.new(mode, (20, 20), value)
                self.assertTrue(_is_blank(uniform, ExtractionLimits()))
        inked = Image.new("I", (20, 20), 60000)
        inked.putpixel((5, 5), 0)
        self.assertEqual(_flatten(inked).getpixel((5, 5)), (0, 0, 0))
        self.assertGreater(_flatten(inked).getpixel((0, 0))[0], 200)

    def test_ledger_bounds_are_utf8_bytes_and_pictures_keep_every_item(self):
        from lifecontext.extraction import compact_ledger, expand_ledger, ledger_items

        cyrillic = [f"chart_not_read:chart:{n}; sheet Бюджет на рік; chart data is in the sheet cells"
                    for n in range(1, 1500, 2)]
        compacted = compact_ledger(cyrillic + [f"ocr_page:{n}; OCR may contain errors" for n in range(2, 400, 2)])
        self.assertTrue(all(len(entry.encode()) <= 1024 for entry in compacted))
        pictures = [f"embedded_image_not_read:picture:{n}; document paragraph {n * 3}" for n in range(1, 1201)]
        compacted = compact_ledger(pictures)
        self.assertLessEqual(len(compacted), 1000)
        items = [item for entry in compacted for item in ledger_items(entry)[1]]
        self.assertEqual(sorted(items), sorted(f"picture:{n}" for n in range(1, 1201)))
        self.assertNotIn("ledger_truncated", " ".join(compacted))
        self.assertEqual(len(expand_ledger(compacted)), 1200)

    def test_alternate_content_falls_back_when_the_choice_lacks_the_element(self):
        from docx import Document
        from docx.oxml import parse_xml

        document = Document()
        paragraph = document.add_paragraph("Body")
        content = (
            '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
            'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:v="urn:schemas-microsoft-com:vml" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            '<mc:Choice Requires="cx1"><w:drawing/></mc:Choice>'
            '<mc:Fallback><w:pict><v:shape><v:imagedata r:id="rIdMissing"/></v:shape></w:pict></mc:Fallback>'
            "</mc:AlternateContent>"
        )
        paragraph.runs[0]._r.append(parse_xml(content))
        buffer = BytesIO()
        document.save(buffer)
        artifact = self.extract(buffer.getvalue(), "a.docx")
        self.assertEqual(artifact.warnings, ("embedded_image_not_read:picture:1; document paragraph 1",))

    def test_sixteen_bit_scans_are_not_mistaken_for_blank(self):
        import struct
        import tempfile
        from pathlib import Path

        from PIL import Image

        from lifecontext.extraction import _binary_extract, _is_blank

        samples = [1000 if 10 <= x < 50 and 10 <= y < 30 else 60000 for y in range(60) for x in range(60)]
        scan = Image.frombytes("I;16", (60, 60), struct.pack("<3600H", *samples))
        self.assertFalse(_is_blank(scan, ExtractionLimits()))
        self.assertTrue(_is_blank(Image.new("I;16", (60, 60), 60000), ExtractionLimits()))
        buffer = BytesIO()
        scan.save(buffer, format="TIFF")
        with tempfile.TemporaryDirectory() as directory, patch(
            "lifecontext.extraction._ocr_image", return_value="Ink"
        ) as ocr:
            result = _binary_extract(buffer.getvalue(), "tiff", ExtractionLimits(), Path(directory), "/mock/tesseract")
        ocr.assert_called_once()
        self.assertEqual(result["warnings"], ["ocr_frame:1; OCR may contain errors"])

    def test_one_handwritten_word_is_not_a_blank_page(self):
        from PIL import Image, ImageDraw

        from lifecontext.extraction import _is_blank

        page = Image.new("RGB", (1275, 1650), "white")
        ImageDraw.Draw(page).text((600, 800), "Notes", fill="black", font_size=28)
        self.assertFalse(_is_blank(page, ExtractionLimits()))
        page = Image.new("RGB", (1275, 1650), "white")
        page.putpixel((10, 10), (0, 0, 0))
        self.assertTrue(_is_blank(page, ExtractionLimits()))

    def test_long_ledgers_are_compacted_without_losing_items(self):
        from lifecontext.extraction import (
            compact_ledger,
            coverage_summary,
            expand_ledger,
        )

        warnings = [f"figure_not_read:{n}; images or drawings on this page were not read" for n in range(1, 3001, 2)]
        warnings += [f"ocr_page:{n}; OCR may contain errors" for n in range(2, 1202, 2)]
        warnings += ["vision_opted_out; this document is never sent to a vision model"]
        compacted = compact_ledger(warnings)
        self.assertLessEqual(len(compacted), 1000)
        self.assertTrue(all(len(entry) <= 1024 for entry in compacted))
        self.assertIn("vision_opted_out; this document is never sent to a vision model", compacted)
        self.assertEqual(sorted(expand_ledger(compacted)), sorted(warnings))
        segments = [{"locator": {"kind": "page", "start": n}} for n in range(1, 3001)]
        full = coverage_summary({"format": "pdf", "segments": segments, "warnings": warnings})
        short = coverage_summary({"format": "pdf", "segments": segments, "warnings": list(compacted)})
        self.assertEqual(full, short)
        self.assertEqual((short["total"], short["ocr"], len(short["partial"])), (3000, 600, 1000))
        self.assertEqual(compact_ledger(warnings[:5]), tuple(warnings[:5]))

    def test_docx_alternate_content_is_read_once(self):
        from docx import Document
        from docx.oxml import parse_xml

        document = Document()
        paragraph = document.add_paragraph("Body")
        box = (
            '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
            'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
            'xmlns:v="urn:schemas-microsoft-com:vml">'
            '<mc:Choice Requires="wps"><w:drawing><wps:txbx><w:txbxContent><w:p><w:r><w:t>Boxed</w:t></w:r></w:p>'
            '</w:txbxContent></wps:txbx></w:drawing></mc:Choice>'
            '<mc:Fallback><w:pict><v:shape><v:textbox><w:txbxContent><w:p><w:r><w:t>Boxed</w:t></w:r></w:p>'
            '</w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback></mc:AlternateContent>'
        )
        paragraph.runs[0]._r.append(parse_xml(box))
        buffer = BytesIO()
        document.save(buffer)
        artifact = self.extract(buffer.getvalue(), "a.docx")
        self.assertEqual(artifact.status, "completed", artifact.reason)
        self.assertEqual(artifact.text, "Body\nBoxed\n")
        self.assertEqual(artifact.segments[1].locator.container, "document/paragraph:1/textbox:1/line:1")

    def test_suspect_text_layer_heuristics(self):
        from lifecontext.extraction import _suspect_text

        limits = ExtractionLimits()
        self.assertFalse(_suspect_text("Plain English text with nothing odd in it at all.", limits))
        self.assertFalse(_suspect_text("Звичайний український текст без жодних проблем тут.", limits))
        self.assertTrue(_suspect_text("Пpивiт cвiт, цe тeкcт iз лaтинcькими літерами всередині.", limits))
        self.assertTrue(_suspect_text("��� broken � glyphs �� everywhere here", limits))
        self.assertTrue(_suspect_text("Ïðèâåò ìèð, ýòî òåêñò â íåïðàâèëüíîé êîäèðîâêå", limits))

    def test_tiff_frames_are_isolated_and_blank_frames_resolved(self):
        from PIL import Image, ImageDraw

        blank = Image.new("RGB", (40, 40), "white")
        inked = Image.new("RGB", (40, 40), "white")
        ImageDraw.Draw(inked).rectangle((5, 5, 30, 30), fill="black")
        buffer = BytesIO()
        blank.save(buffer, format="TIFF", save_all=True, append_images=[inked, inked.copy()])
        with patch("lifecontext.extraction._ocr_image", side_effect=["Frame text", ""]):
            import tempfile
            from pathlib import Path

            from lifecontext.extraction import _binary_extract

            with tempfile.TemporaryDirectory() as directory:
                result = _binary_extract(
                    buffer.getvalue(), "tiff", ExtractionLimits(), Path(directory), "/mock/tesseract"
                )
        self.assertEqual(result["status"], "completed")
        self.assertEqual([s["locator"]["start"] for s in result["segments"]], [2])
        self.assertEqual(
            result["warnings"],
            [
                "blank_frame:1; no visible content",
                "ocr_frame:2; OCR may contain errors",
                "ocr_no_text_frame:3; frame content unresolved",
            ],
        )

    def test_routed_images_are_budgeted_files_listed_in_the_result(self):
        import tempfile
        from pathlib import Path

        from lifecontext.extraction import _binary_extract

        with tempfile.TemporaryDirectory() as directory, patch(
            "lifecontext.extraction._ocr_image", return_value="Scanned"
        ):
            root = Path(directory)
            result = _binary_extract(
                self.mixed_pdf(), "pdf", ExtractionLimits(), root, "/mock/tesseract", 5
            )
            [entry] = result["images"]
            self.assertEqual(
                {k: entry[k] for k in ("item", "mode", "kind", "index", "after", "replaces")},
                {"item": "page:2", "mode": "page", "kind": "page", "index": 2, "after": 1, "replaces": 1},
            )
            self.assertEqual(entry["warnings"], ["ocr_page:2; OCR may contain errors"])
            self.assertTrue((root / entry["file"]).read_bytes().startswith(b"\xff\xd8\xff"))
        with tempfile.TemporaryDirectory() as directory, patch(
            "lifecontext.extraction._ocr_image", return_value=""
        ):
            result = _binary_extract(
                self.pdf(drawing=True, pages=2), "pdf", ExtractionLimits(), Path(directory), "/mock/tesseract", 1
            )
            # Nothing was read, but a vision reader may still resolve the routed page.
            self.assertEqual((result["status"], result["reason"]), ("failed", "no_extractable_text"))
            self.assertEqual([e["item"] for e in result["images"]], ["page:1"])
            self.assertIn("vision_skipped:page:2; vision page budget reached", result["warnings"])
        # Without a vision budget no image is ever written.
        with tempfile.TemporaryDirectory() as directory, patch(
            "lifecontext.extraction._ocr_image", return_value=""
        ):
            result = _binary_extract(
                self.mixed_pdf(), "pdf", ExtractionLimits(), Path(directory), "/mock/tesseract"
            )
            self.assertEqual(result["images"], [])
            self.assertEqual([p.name for p in Path(directory).iterdir()], [])

    def test_routed_images_survive_the_worker_directory(self):
        from lifecontext.extraction import extract_document_with_images

        artifact, routed = extract_document_with_images(
            self.mixed_pdf(), "a.pdf", source_id="document-1", vision_items=3
        )
        self.assertEqual(artifact.status, "completed", artifact.reason)
        [image] = routed
        self.assertEqual((image.item, image.after, image.replaces), ("page:2", 1, None))
        self.assertEqual(image.resolves, ("ocr_unavailable_page:2; page content unresolved",))
        self.assertTrue(image.data.startswith(b"\xff\xd8\xff"))
        artifact, routed = extract_document_with_images(self.mixed_pdf(), "a.pdf", source_id="document-1")
        self.assertEqual(routed, ())

    def test_coverage_summary_accounts_for_every_page(self):
        from lifecontext.extraction import coverage_summary

        payload = {
            "format": "pdf",
            "segments": [
                {"locator": {"kind": "page", "start": 1}},
                {"locator": {"kind": "page", "start": 2}},
                {"locator": {"kind": "page", "start": 5}},
            ],
            "warnings": [
                "ocr_page:2; OCR may contain errors",
                "blank_page:3; no visible content",
                "ocr_no_text_page:4; page content unresolved",
                "figure_not_read:1; images or drawings on this page were not read",
                "vision_page:5; read by model-x; verify against the original page",
            ],
        }
        self.assertEqual(
            coverage_summary(payload),
            {
                "unit": "page", "total": 5, "resolved": 4, "text": 1, "ocr": 1, "blank": 1,
                "vision": 1, "unresolved": ["page:4"], "partial": ["figure_not_read:1"],
                "complete": False,
            },
        )
        office = coverage_summary(
            {"format": "docx", "segments": [{"locator": {"kind": "paragraph", "start": 1}}],
             "warnings": ["embedded_image_not_read:picture:1; document paragraph 2"]}
        )
        self.assertEqual(
            (office["unit"], office["total"], office["partial"], office["complete"]),
            (None, None, ["embedded_image_not_read:picture:1"], False),
        )

    def test_docx_headers_footers_notes_text_boxes_and_pictures(self):
        from docx import Document
        from PIL import Image

        document = Document()
        document.add_paragraph("Body")
        picture = BytesIO()
        Image.new("RGB", (8, 8), "red").save(picture, format="PNG")
        document.add_picture(picture)
        section = document.sections[0]
        section.header.paragraphs[0].text = "Running head"
        section.footer.paragraphs[0].text = "Page footer"
        buffer = BytesIO()
        document.save(buffer)
        # python-docx cannot write footnotes; add the part the way Word stores it.
        source = ZipFile(BytesIO(buffer.getvalue()))
        rebuilt = BytesIO()
        with ZipFile(rebuilt, "w") as target:
            for entry in source.infolist():
                target.writestr(entry, source.read(entry))
            target.writestr(
                "word/footnotes.xml",
                b'<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                b'<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:t>---</w:t></w:r></w:p></w:footnote>'
                b'<w:footnote w:id="1"><w:p><w:r><w:t>A footnote</w:t></w:r></w:p></w:footnote>'
                b"</w:footnotes>",
            )
        artifact = self.extract(rebuilt.getvalue(), "a.docx")
        self.assertEqual(artifact.status, "completed", artifact.reason)
        self.assertEqual(artifact.text, "Body\nRunning head\nPage footer\nA footnote\n")
        self.assertEqual(
            [s.locator.container for s in artifact.segments],
            ["document", "section:1/header", "section:1/footer", "footnote:1"],
        )
        self.assertEqual(
            artifact.warnings, ("embedded_image_not_read:picture:1; document paragraph 2",)
        )

    def test_pptx_notes_charts_and_pictures(self):
        from PIL import Image
        from pptx import Presentation
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.util import Inches

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        data = CategoryChartData()
        data.categories = ["Q1", "Q2"]
        data.add_series("Revenue", (10, 20.5))
        slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED, 0, 0, Inches(4), Inches(3), data
        )
        picture = BytesIO()
        Image.new("RGB", (8, 8), "blue").save(picture, format="PNG")
        slide.shapes.add_picture(picture, 0, 0)
        slide.notes_slide.notes_text_frame.text = "Speaker notes"
        buffer = BytesIO()
        presentation.save(buffer)
        artifact = self.extract(buffer.getvalue(), "a.pptx")
        self.assertEqual(artifact.status, "completed", artifact.reason)
        self.assertEqual(
            artifact.text,
            "| Category | Revenue |\n| --- | --- |\n| Q1 | 10.0 |\n| Q2 | 20.5 |\nSpeaker notes\n",
        )
        self.assertTrue(artifact.segments[0].locator.container.endswith("/chart"))
        self.assertEqual(artifact.segments[1].locator.container, "slide:1/notes")
        self.assertEqual(artifact.warnings, ("embedded_image_not_read:picture:1; slide 1",))

    def test_xlsx_charts_and_pictures_enter_the_ledger(self):
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference
        from openpyxl.drawing.image import Image as SheetImage
        from PIL import Image

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Budget"
        for row in (("Month", "Spend"), ("Jan", 3), ("Feb", 5)):
            sheet.append(row)
        chart = BarChart()
        chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
        sheet.add_chart(chart, "D2")
        picture = BytesIO()
        Image.new("RGB", (8, 8), "green").save(picture, format="PNG")
        sheet.add_image(SheetImage(picture), "D20")
        buffer = BytesIO()
        workbook.save(buffer)
        artifact = self.extract(buffer.getvalue(), "a.xlsx")
        self.assertEqual(artifact.status, "completed", artifact.reason)
        self.assertEqual(
            artifact.warnings[1:],
            (
                "chart_not_read:chart:1; sheet Budget; chart data is in the sheet cells",
                "embedded_image_not_read:picture:1; sheet Budget",
            ),
        )

    def test_ocr_languages_only_selects_available_data(self):
        from lifecontext.extraction import _ocr_languages
        _ocr_languages.cache_clear()
        with patch("lifecontext.extraction.subprocess.check_output", return_value=b"List of available languages:\neng\nosd\nrus\n"):
            self.assertEqual(_ocr_languages("/test/tesseract"), "eng+rus")
        _ocr_languages.cache_clear()

    def test_pdfium_real_render_with_mock_ocr_and_pixel_limit(self):
        import tempfile
        from pathlib import Path

        from lifecontext.extraction import _binary_extract

        with tempfile.TemporaryDirectory() as directory:
            with patch(
                "lifecontext.extraction._ocr_image", return_value="Scanned words"
            ) as ocr:
                result = _binary_extract(
                    self.pdf(drawing=True),
                    "pdf",
                    ExtractionLimits(),
                    Path(directory),
                    "/mock/tesseract",
                )
                self.assertEqual(result["segments"][0]["text"], "Scanned words\n")
                self.assertEqual(result["segments"][0]["locator"]["kind"], "page")
                ocr.assert_called_once()
            # An oversized page is one unresolved page, not a failed document.
            result = _binary_extract(
                self.mixed_pdf(),
                "pdf",
                ExtractionLimits(max_pixels=1),
                Path(directory),
                "/mock/tesseract",
            )
            self.assertEqual(
                result["warnings"],
                ["ocr_failed_page:2; pixel_limit; page content unresolved"],
            )

    def test_missing_dependency_explicit_block(self):
        import builtins
        import tempfile
        from pathlib import Path

        from lifecontext.extraction import _worker_main

        original_import = builtins.__import__

        def missing(name, *args, **kwargs):
            if name == "pypdf":
                raise ImportError("missing")
            return original_import(name, *args, **kwargs)

        with tempfile.TemporaryDirectory() as directory:
            import json

            root = Path(directory)
            (root / "original").write_bytes(self.pdf())
            (root / "request.json").write_text(
                json.dumps(
                    {
                        "format": "pdf",
                        "limits": dataclasses.asdict(ExtractionLimits()),
                        "tesseract": None,
                    }
                )
            )
            with (
                patch("resource.setrlimit"),
                patch("builtins.__import__", side_effect=missing),
            ):
                _worker_main(directory)
            result = json.loads((root / "result.json").read_text())
            self.assertEqual(result["status"], "blocked")
            self.assertEqual(result["reason"], "extraction_dependency_unavailable")
            self.assertEqual(result["segments"], [])

    def test_real_ocr_runtime(self):
        import shutil

        from PIL import Image, ImageDraw

        executable = shutil.which("tesseract")
        if not executable:
            self.skipTest(
                "local Tesseract runtime not installed; real OCR acceptance pending"
            )
        image = Image.new("RGB", (400, 100), "white")
        ImageDraw.Draw(image).text(
            (20, 20), "HELLO DOCUMENT", fill="black", font_size=28
        )
        buffer = BytesIO()
        image.save(buffer, format="PNG")
        artifact = self.extract(buffer.getvalue(), "a.png", tesseract_path=executable)
        self.assertEqual(artifact.status, "completed", artifact.reason)
        self.assertIn("HELLO", artifact.text)


if __name__ == "__main__":
    unittest.main()
